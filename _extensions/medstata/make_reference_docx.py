"""Build reference.docx: a Word reference document that mirrors the medstata typst template.

Build:  python3 make_reference_docx.py            (writes reference.docx next to this file)
Render: quarto render report.qmd --to docx -M reference-doc:<path>/reference.docx
Finish: python3 make_reference_docx.py --fix report.docx

--fix does what a reference doc cannot: fills the header/footer fields' shown values
(short title, version, author; pandoc stores YAML keys as document properties, which Word
only re-reads on F9 or print), sets table cell text to the sans "Table Text" style (pandoc
puts cells in "Compact", shared with tight lists), and drops the duplicate centring pandoc
writes into table cells and the inline alignment on captions.
"""
import re
import subprocess
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).parent
OUT = HERE / "reference.docx"

# Tokens from typst-template.typ
INK = RGBColor(0x1A, 0x1A, 0x1A)
NAVY = RGBColor(0x1A, 0x47, 0x6F)
MUTED = RGBColor(0x5B, 0x67, 0x6C)
HEAD = RGBColor(0x4D, 0x6A, 0x61)  # accent_color #6e8e84 darkened 20%, running head
RULE, TINT, CODE_TINT, ACCENT = "C9D3DD", "EEF2F6", "F2F2EF", "6E8E84"
SERIF, SANS, MONO = "Libertinus Serif", "IBM Plex Sans", "IBM Plex Mono"


def el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(f"w:{k}"), str(v))
    return e


def replace_child(parent, child):
    for old in parent.findall(child.tag):
        parent.remove(old)
    parent.append(child)


def style(doc, name, font=None, size=None, color=None, bold=None, italic=None,
          align=None, before=None, after=None, line=None, keep_next=None):
    s = doc.styles[name]
    f = s.font
    if font:
        f.name = font
        rfonts = s.element.get_or_add_rPr().get_or_add_rFonts()
        for a in ("ascii", "hAnsi", "cs", "eastAsia"):
            rfonts.set(qn(f"w:{a}"), font)
        for a in ("asciiTheme", "hAnsiTheme", "cstheme", "eastAsiaTheme"):
            rfonts.attrib.pop(qn(f"w:{a}"), None)
    if size is not None:
        f.size = Pt(size)
    if color is not None:
        f.color.rgb = color
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    pf = getattr(s, "paragraph_format", None)
    if pf is not None:
        if align is not None:
            pf.alignment = align
        if before is not None:
            pf.space_before = Pt(before)
        if after is not None:
            pf.space_after = Pt(after)
        if line is not None:
            pf.line_spacing = line
        if keep_next is not None:
            pf.keep_with_next = keep_next
    return s


def para_border(s, side, sz, color, space=4):
    ppr = s.element.get_or_add_pPr()
    bdr = ppr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = el("w:pBdr")
        ppr.append(bdr)
    replace_child(bdr, el(f"w:{side}", val="single", sz=sz, space=space, color=color))


def shade(parent, fill):
    replace_child(parent, el("w:shd", val="clear", color="auto", fill=fill))


def field(p, instr, placeholder):
    """Complex field (begin / instr / separate / cached result / end) in paragraph p."""
    def run(child):
        r = p.add_run()
        r._r.append(child)
        return r
    run(el("w:fldChar", fldCharType="begin"))
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = f" {instr} "
    run(instr_el)
    run(el("w:fldChar", fldCharType="separate"))
    p.add_run(placeholder)
    run(el("w:fldChar", fldCharType="end"))


def main():
    default = HERE / "_pandoc_default_reference.docx"
    subprocess.run(["quarto", "pandoc", "-o", str(default), "--print-default-data-file",
                    "reference.docx"], check=True)
    doc = Document(default)
    default.unlink()

    # Document default font: no theme fonts left for Word to substitute
    dflt = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    fonts = el("w:rFonts", ascii=SERIF, hAnsi=SERIF, cs=SERIF, eastAsia=SERIF)
    replace_child(dflt, fonts)

    # Body
    style(doc, "Normal", SERIF, 10, INK)
    for name in ("Body Text", "First Paragraph"):
        style(doc, name, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=9, line=1.15)
    style(doc, "Compact", before=0, after=2, line=1.1)
    style(doc, "Block Text", size=9.5, before=6, after=9)
    para_border(doc.styles["Block Text"], "left", 18, "1A476F", 8)
    style(doc, "Footnote Text", size=8.5)
    style(doc, "Hyperlink", color=NAVY)
    doc.styles["Hyperlink"].font.underline = False

    # Title block (the typst cover is a full dark page; Word gets a quiet equivalent)
    style(doc, "Title", SERIF, 22, NAVY, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT,
          before=120, after=12)
    doc.styles["Title"].font.all_caps = True
    para_border(doc.styles["Title"], "bottom", 6, ACCENT, 6)
    style(doc, "Subtitle", SERIF, 14, INK, bold=False, italic=False,
          align=WD_ALIGN_PARAGRAPH.LEFT, before=6, after=36)
    doc.styles["Subtitle"].font.all_caps = False
    bdr = el("w:pBdr")  # Subtitle is based on Title: switch off the inherited rule
    bdr.append(el("w:bottom", val="nil"))
    replace_child(doc.styles["Subtitle"].element.get_or_add_pPr(), bdr)
    for name in ("Author", "Date"):
        style(doc, name, SANS, 10, MUTED, align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=2)
        doc.styles[name].font.all_caps = False
    style(doc, "Abstract Title", SANS, 11, NAVY, bold=True)
    style(doc, "Abstract", size=10)

    # Headings: IBM Plex Sans, navy (typst: 15 / 11 / 9.5 pt semibold)
    for lvl, size, before, after in ((1, 15, 20, 8), (2, 11, 14, 5), (3, 9.5, 10, 4)):
        style(doc, f"Heading {lvl}", SANS, size, NAVY, bold=True, italic=False,
              align=WD_ALIGN_PARAGRAPH.LEFT, before=before, after=after, keep_next=True)
    for lvl in range(4, 10):
        style(doc, f"Heading {lvl}", SANS, 10, INK, bold=True, italic=False,
              before=8, after=3, keep_next=True)
    style(doc, "TOC Heading", SANS, 15, NAVY, bold=True)
    # Linked character styles carry their own (theme) fonts: align them with their paragraph style
    for name, font in (("Title Char", SERIF), ("Subtitle Char", SERIF),
                       ("Heading 1 Char", SANS), ("Heading 2 Char", SANS), ("Heading 3 Char", SANS)):
        style(doc, name, font)

    # Captions: original medstata style, body font, black, centred
    for name in ("Caption", "Table Caption", "Image Caption"):
        style(doc, name, SERIF, 10, INK, bold=False, italic=False,
              align=WD_ALIGN_PARAGRAPH.CENTER, before=10, after=5)
    doc.styles["Table Caption"].paragraph_format.keep_with_next = True
    style(doc, "Captioned Figure", align=WD_ALIGN_PARAGRAPH.CENTER, keep_next=True)
    style(doc, "Figure", align=WD_ALIGN_PARAGRAPH.CENTER)

    # Code (search strings): IBM Plex Mono on a barely-off-page panel, navy left rule
    style(doc, "Verbatim Char", MONO, 8)
    if "Source Code" not in [s.name for s in doc.styles]:
        sc = doc.styles.add_style("Source Code", WD_STYLE_TYPE.PARAGRAPH)
        sc.base_style = doc.styles["Normal"]
        sc.element.set(qn("w:customStyle"), "1")
    style(doc, "Source Code", MONO, 8, INK, align=WD_ALIGN_PARAGRAPH.LEFT,
          before=2, after=8, line=1.0)
    shade(doc.styles["Source Code"].element.get_or_add_pPr(), CODE_TINT)
    para_border(doc.styles["Source Code"], "left", 12, "1A476F", 6)

    # References
    style(doc, "Bibliography", SERIF, 9.5, INK, align=WD_ALIGN_PARAGRAPH.LEFT, after=4)
    doc.styles["Bibliography"].paragraph_format.left_indent = Cm(0.8)
    doc.styles["Bibliography"].paragraph_format.first_line_indent = Cm(-0.8)

    # Cell text (applied by --fix): Plex Sans 8.5 pt, single-spaced, left
    tt = doc.styles.add_style("Table Text", WD_STYLE_TYPE.PARAGRAPH)
    tt.base_style = doc.styles["Normal"]
    style(doc, "Table Text", SANS, 8.5, INK, align=WD_ALIGN_PARAGRAPH.LEFT,
          before=0, after=0, line=1.0)

    # Table style (pandoc names it "Table"): Plex Sans 8.5 pt, tinted header row with navy
    # bold text, navy rule above the header and under the table, hairline rules between rows.
    t = doc.styles["Table"]
    t.font.name = SANS
    rf = t.element.get_or_add_rPr().get_or_add_rFonts()
    for a in ("ascii", "hAnsi", "cs"):
        rf.set(qn(f"w:{a}"), SANS)
    t.font.size = Pt(8.5)
    tpr = t.element.find(qn("w:tblPr"))
    if tpr is None:
        tpr = el("w:tblPr")
        t.element.append(tpr)
    borders = el("w:tblBorders")
    borders.append(el("w:top", val="single", sz=6, space=0, color="1A476F"))
    borders.append(el("w:bottom", val="single", sz=6, space=0, color="1A476F"))
    borders.append(el("w:insideH", val="single", sz=3, space=0, color=RULE))
    for side in ("left", "right", "insideV"):
        borders.append(el(f"w:{side}", val="nil"))
    replace_child(tpr, borders)
    mar = el("w:tblCellMar")
    for side, w in (("top", 60), ("bottom", 60), ("left", 100), ("right", 100)):
        mar.append(el(f"w:{side}", w=w, type="dxa"))
    replace_child(tpr, mar)
    for old in t.element.findall(qn("w:tblStylePr")):
        t.element.remove(old)
    first = el("w:tblStylePr", type="firstRow")
    frpr = el("w:rPr")
    frpr.append(el("w:b"))
    frpr.append(el("w:color", val="1A476F"))
    first.append(frpr)
    tcpr = el("w:tcPr")
    tcb = el("w:tcBorders")
    tcb.append(el("w:top", val="single", sz=6, space=0, color="1A476F"))
    tcb.append(el("w:bottom", val="single", sz=6, space=0, color="1A476F"))
    tcpr.append(tcb)
    shade(tcpr, TINT)
    first.append(tcpr)
    first_trpr = el("w:trPr")
    first_trpr.append(el("w:tblHeader"))
    first.insert(0, el("w:pPr"))
    first.append(first_trpr)
    t.element.append(first)
    ppr = t.element.get_or_add_pPr()
    replace_child(ppr, el("w:spacing", before=0, after=0, line=240, lineRule="auto"))
    replace_child(ppr, el("w:jc", val="left"))

    # Page: A4, 2.5 cm margins, running header and footer as in the typst template
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, side, Cm(2.5))
    sec.header_distance = sec.footer_distance = Cm(1.2)

    hp = sec.header.paragraphs[0]
    # Left: running chapter (STYLEREF picks up the current Heading 1); right: short title | version
    field(hp, 'STYLEREF "Heading 1" \\* MERGEFORMAT', "Chapter")
    # positional tab relative to the margin: right-aligned on portrait and landscape pages
    hp.add_run()._r.append(el("w:ptab", relativeTo="margin", alignment="right", leader="none"))
    field(hp, "DOCPROPERTY short-title \\* Upper", "SHORT TITLE")
    hp.add_run("   |   ")
    hp.add_run("v")
    field(hp, "DOCPROPERTY version", "VERSION")
    for r in hp.runs:
        r.font.name, r.font.size, r.font.color.rgb = SERIF, Pt(7.5), HEAD
    pbdr = el("w:pBdr")
    pbdr.append(el("w:bottom", val="single", sz=4, space=3, color=ACCENT))
    hp._p.get_or_add_pPr().append(pbdr)

    fp = sec.footer.paragraphs[0]
    field(fp, "AUTHOR", "AUTHOR")
    fp.add_run()._r.append(el("w:ptab", relativeTo="margin", alignment="right", leader="none"))
    fp.add_run("Page ")
    field(fp, "PAGE", "1")
    fp.add_run(" of ")
    field(fp, "NUMPAGES", "1")
    for r in fp.runs:
        r.font.name, r.font.size, r.font.color.rgb = SERIF, Pt(7.5), MUTED

    doc.save(OUT)
    print(f"wrote {OUT}")


def fix(path):
    """Post-render fixes on a pandoc/quarto docx made with this reference doc."""
    with zipfile.ZipFile(path) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    core = parts["docProps/core.xml"].decode()
    custom = parts.get("docProps/custom.xml", b"").decode()

    def prop(name):
        m = re.search(rf'name="{re.escape(name)}"><vt:lpwstr>([^<]*)<', custom)
        return m.group(1) if m else ""

    creator = re.search(r"<dc:creator>([^<]*)<", core)
    title = re.search(r"<dc:title>([^<]*)<", core)
    values = {
        "SHORT TITLE": (prop("short-title") or (title.group(1) if title else "")).upper(),
        "VERSION": prop("version"),
        "AUTHOR": creator.group(1) if creator else "",
    }
    for name in [n for n in parts if re.match(r"word/(header|footer)\d*\.xml", n)]:
        xml = parts[name].decode()
        for placeholder, value in values.items():
            xml = xml.replace(f">{placeholder}</w:t>", f">{value}</w:t>")
        parts[name] = xml.encode()

    doc = parts["word/document.xml"].decode()

    # Section breaks inserted for landscape pages carry no header, footer or margins: copy them
    # from the final section so every page gets the running header and footer
    final = re.findall(r"<w:sectPr>(?:(?!<w:sectPr>).)*?</w:sectPr>", doc, re.S)[-1]
    refs = "".join(re.findall(r"<w:(?:header|footer)Reference [^>]*/>", final))
    mar = re.search(r"<w:pgMar [^>]*/>", final).group(0)
    portrait = re.search(r"<w:pgSz [^>]*/>", final).group(0)

    def section(m):
        body = m.group(1)
        if "headerReference" in body:
            return m.group(0)
        size = re.search(r"<w:pgSz [^>]*/>", body)
        return f"<w:sectPr>{refs}{size.group(0) if size else portrait}{mar}</w:sectPr>"

    doc = re.sub(r"<w:sectPr>((?:(?!</w:sectPr>).)*)</w:sectPr>", section, doc, flags=re.S)

    def cells(m):
        t = m.group(0).replace('<w:pStyle w:val="Compact" />', '<w:pStyle w:val="TableText" />')
        # pandoc writes the column alignment then a stray centring: keep the first
        return re.sub(r'(<w:jc w:val="\w+" />)<w:jc w:val="center" />', r"\1", t)

    doc = re.sub(r"<w:tbl>.*?</w:tbl>", cells, doc, flags=re.S)
    # quarto wraps each captioned table in a one-cell layout table: no rules on that one
    doc = re.sub(r'(<w:tblStyle w:val="Table" /><w:tblW [^>]*/>)(<w:tblLayout [^>]*/><w:tblLook w:firstRow="0"[^>]*w:val="0000" />)',
                 r'\1<w:tblBorders><w:top w:val="nil"/><w:left w:val="nil"/><w:bottom w:val="nil"/>'
                 r'<w:right w:val="nil"/><w:insideH w:val="nil"/><w:insideV w:val="nil"/></w:tblBorders>\2', doc)
    # captions: let the caption style centre them
    doc = re.sub(r'<w:p><w:pPr><w:jc w:val="center" /></w:pPr><w:pPr>\s*<w:jc w:val="left"/>',
                 "<w:p><w:pPr>", doc)
    parts["word/document.xml"] = doc.encode()

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n, data in parts.items():
            z.writestr(n, data)
    print(f"fixed {path}")


if __name__ == "__main__":
    if sys.argv[1:2] == ["--fix"]:
        for p in sys.argv[2:]:
            fix(p)
    else:
        main()
